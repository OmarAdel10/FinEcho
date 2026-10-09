from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path('deliverables/FinEcho_Project_Proposal.docx')


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_font(run, size=10.5, bold=False, color='212121'):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def cell_text(cell, text, bold=False, color='212121', align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_font(r, 9.5, bold, color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margin(cell)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, 15 if level == 1 else 11.5, True, '000000')
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_font(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        shade(cell, '003C33')
        cell_text(cell, header, True, 'FFFFFF', WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].width = Inches(widths[i])
            if index % 2 == 1:
                shade(cells[i], 'F4F7F6')
            cell_text(cells[i], value, align=WD_ALIGN_PARAGRAPH.CENTER if i in (0, 1) else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    styles['Normal'].font.size = Pt(10.5)

    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run('FinEcho Mobile Application Project Proposal')
    set_font(r, 24, True, '000000')

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run('Mobile Development and Advanced Software Engineering')
    set_font(r, 11, False, '616161')

    info = add_table(doc, ['Course / Section', 'Project delivery deadline'], [
        ['[Insert course name and section]', '[Insert instructor deadline]'],
    ], [3.25, 3.25])

    add_heading(doc, 'Team Members', 1)
    add_body(doc, 'Replace the bracketed placeholders below with the names and university IDs of the four project members before submission.')
    add_table(doc, ['No.', 'Student full name', 'University ID', 'Primary responsibility'], [
        ['1', '[Student 1 name]', '[Student 1 ID]', 'Architecture and encrypted local data layer'],
        ['2', '[Student 2 name]', '[Student 2 ID]', 'Offline speech and categorization features'],
        ['3', '[Student 3 name]', '[Student 3 ID]', 'Android SMS integration and local notifications'],
        ['4', '[Student 4 name]', '[Student 4 ID]', 'Flutter UI, accessibility, analytics, and testing'],
    ], [0.4, 1.75, 1.15, 3.2])

    add_heading(doc, 'Project Description and Summary', 1)
    add_body(doc, 'FinEcho is an offline-first personal finance tracking application developed with Flutter. It is designed to help users record income and expenses, organize multiple accounts, set category budgets, and understand spending patterns without sending financial records to a cloud service.')
    add_body(doc, 'The application combines a straightforward manual ledger with optional on-device automation. Users can add a transaction by voice using offline speech-to-text, while supported Android devices can locally parse transactional SMS messages after the user grants permission. A transparent categorization process proposes categories and always allows the user to review or correct the result.')
    add_body(doc, 'The planned interface includes an animated splash screen, three onboarding screens, a local security gate, Home, Accounts, Statistics, and Settings screens. Encrypted SQLite storage, biometric/PIN access, local notifications, budgets, and data exports form the privacy and reliability foundation of the app. SMS automation is intentionally scoped to Android because iOS does not expose a general SMS inbox interface.')

    add_heading(doc, 'Key Objectives', 1)
    for item in [
        'Deliver a simple and intuitive personal-finance workflow that works without internet access.',
        'Use encrypted local storage and device security mechanisms to protect financial records.',
        'Provide optional on-device voice entry, Android SMS parsing, card linking, and automatic categorization.',
        'Help users act on their data through budgets, statistics, spending-pattern insights, and local reminders.',
        'Maintain quality through automated formatting, analysis, and test checks in continuous integration.',
    ]:
        add_bullet(doc, item)

    add_heading(doc, 'Delivery Timeline', 1)
    add_body(doc, 'The project is planned as a 16-week delivery cycle. The team should align Week 1 with the course start date and confirm the final submission date with the instructor.')
    add_table(doc, ['Sprint', 'Weeks', 'Milestone and deliverables'], [
        ['1', '1-2', 'Project setup, architecture, encrypted SQLite schema, baseline CI, and initial navigation.'],
        ['2', '3-5', 'Accounts, manual transactions, Android SMS parsing prototype, card linking, and local notifications.'],
        ['3', '6-8', 'Offline speech-to-text, voice transaction review, and entity extraction.'],
        ['4', '9-11', 'Auto-categorization, budgets, recurring-payment detection, and scheduled reminders.'],
        ['5', '12-14', 'Cohere-inspired Flutter UI, statistics, accessibility refinement, and integration testing.'],
        ['6', '15-16', 'Performance profiling, CI/CD release preparation, documentation, final testing, and submission.'],
    ], [0.55, 0.75, 5.95])

    add_heading(doc, 'Delivery Evidence', 1)
    add_body(doc, 'The final submission will include the Flutter source code, technical specifications, user-flow and architecture diagrams, automated test results, CI/CD configuration, and a short demonstration of the working application on an appropriate emulator or device.')

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run('FinEcho Project Proposal')
    set_font(fr, 8.5, False, '616161')

    doc.core_properties.title = 'FinEcho Mobile Application Project Proposal'
    doc.core_properties.subject = 'University mobile development project proposal'
    doc.core_properties.author = 'FinEcho Project Team'
    doc.save(OUT)


if __name__ == '__main__':
    main()
